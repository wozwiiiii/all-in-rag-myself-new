"""分块逻辑测试：标题路径保留、超长切分、docx 提取、扫描 PDF 的 OCR 兜底。"""

from pathlib import Path

from app.rag import chunking
from app.rag.chunking import Block, Chunk, chunk_blocks, extract_docx


def test_heading_path_preserved():
    blocks = [
        Block(level=1, text="第一章 总则"),
        Block(level=0, text="为了加强安全管理，制定本规范。"),
        Block(level=2, text="适用范围"),
        Block(level=0, text="适用于危险化学品企业。"),
    ]
    chunks = chunk_blocks(blocks, source="a.docx")
    assert len(chunks) == 2
    assert chunks[0].heading_path == "第一章 总则"
    assert chunks[1].heading_path == "第一章 总则 > 适用范围"
    assert all(c.source == "a.docx" for c in chunks)


def test_long_section_split_with_overlap():
    paras = [f"第{i}段落内容，" + "字" * 200 for i in range(10)]
    blocks = [Block(level=1, text="标题"), *[Block(level=0, text=p) for p in paras]]
    chunks = chunk_blocks(blocks, source="b.pdf")
    assert len(chunks) > 1
    # 均不超过上限（拼接单段硬切场景下有 overlap 冗余，允许少量超出）
    for c in chunks:
        assert len(c.text) <= 1200 + 200


def test_empty_blocks_no_chunks():
    assert chunk_blocks([], source="x") == []


def test_extract_docx(tmp_path: Path):
    import docx

    doc = docx.Document()
    doc.add_heading("规章标题", level=1)
    doc.add_paragraph("正文内容一。")
    doc.add_heading("细则", level=2)
    doc.add_paragraph("正文内容二。")
    path = tmp_path / "sample.docx"
    doc.save(str(path))

    blocks = extract_docx(path)
    assert (blocks[0].level, blocks[0].text) == (1, "规章标题")
    assert any(b.level == 0 and "正文内容一" in b.text for b in blocks)

    chunks = chunk_blocks(blocks, source="sample.docx")
    assert len(chunks) == 2
    assert all(isinstance(c, Chunk) for c in chunks)


def test_extract_pdf_ocr_fallback(tmp_path: Path, monkeypatch):
    """无文本层的页应触发 OCR；OCR 结果进入 blocks。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)  # 空白页：extract_text 为空
    path = tmp_path / "scanned.pdf"
    with open(path, "wb") as f:
        writer.write(f)

    monkeypatch.setattr(chunking, "_ocr_pdf_page", lambda p, i, **kw: "OCR识别出的文本")
    blocks = chunking.extract_pdf(path)
    assert len(blocks) == 1
    assert blocks[0].text == "OCR识别出的文本"


def test_extract_pdf_ocr_deps_missing(tmp_path: Path, monkeypatch):
    """OCR 依赖缺失时返回空串，不抛异常（优雅降级）。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    path = tmp_path / "scanned.pdf"
    with open(path, "wb") as f:
        writer.write(f)

    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name in ("pypdfium2", "rapidocr_onnxruntime"):
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    assert chunking._ocr_pdf_page(path, 0) == ""
    # 降级后 extract_pdf 不产生该页 block
    blocks = chunking.extract_pdf(path)
    assert blocks == []
