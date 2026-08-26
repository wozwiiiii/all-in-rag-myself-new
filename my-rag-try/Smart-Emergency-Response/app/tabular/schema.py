"""表结构元数据：官方 docx 解析 + 按需生成所选表的紧凑画像。

设计约束：结构说明 docx 是无表边界的大字段表（约 2.3 万字符），全量注入
system prompt 会导致 7B 模型推理超时。因此：
1. 解析为全局"列名 -> 描述"字典（一次性，进程内缓存）；
2. 代理先做表选择，再只为所选表生成紧凑 schema（列名=d描述 + dtype + 1 行示例）。
"""

from __future__ import annotations

import re

from app.core.config import Settings
from app.core.logging import get_logger

logger = get_logger(__name__)

SCHEMA_DOC_FILES = ["数据库危化数据表结构.docx", "数据表介绍.docx"]


def docx_to_markdown(path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style is not None else ""
        if style.startswith("Heading"):
            digits = "".join(c for c in style.split()[-1] if c.isdigit())
            level = min(int(digits) if digits else 1, 4)
            parts.append("#" * level + " " + text)
        else:
            parts.append(text)
    for table in doc.tables:
        rows = []
        for i, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
            if i == 0:
                rows.append("|" + "---|" * len(cells))
        parts.append("\n".join(rows))
    return "\n\n".join(parts)


_DOCS_CACHE: dict[str, str] = {}


def load_schema_docs(settings: Settings) -> str:
    key = str(settings.table_data_dir)
    if key in _DOCS_CACHE:
        return _DOCS_CACHE[key]
    parts: list[str] = []
    for name in SCHEMA_DOC_FILES:
        path = settings.table_data_dir / name
        if not path.exists():
            logger.warning("schema_doc_missing", extra={"file": name})
            continue
        parts.append(f"## {name}\n\n{docx_to_markdown(path)}")
    if not parts:
        raise RuntimeError(f"未找到表结构说明文档: {SCHEMA_DOC_FILES}")
    docs = "\n\n".join(parts)
    _DOCS_CACHE[key] = docs
    return docs


_COLUMN_CACHE: dict[str, dict[str, str]] = {}


def column_dict(settings: Settings) -> dict[str, str]:
    """从结构 docx 的字段表解析 列名 -> 中文描述（后表覆盖前表，同名列语义基本一致）。"""
    key = str(settings.table_data_dir)
    if key in _COLUMN_CACHE:
        return _COLUMN_CACHE[key]
    docs = load_schema_docs(settings)
    result: dict[str, str] = {}
    for line in docs.split("\n"):
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        # 字段表格式: 列名 | 描述 | 字符类型 | 长度
        if len(cells) >= 2 and cells[0] not in ("列名", "") and cells[0] != "---":
            if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", cells[0]):
                result[cells[0]] = cells[1]
    logger.info("column_dict_built", extra={"columns": len(result)})
    _COLUMN_CACHE[key] = result
    return result


_INTROS_CACHE: dict[str, dict[str, str]] = {}


def _physical_name_map(settings: Settings) -> dict[str, str]:
    """从《表结构》docx 开头的表清单提取 中文名 -> 物理表名（如 动火作业 -> t_dhzy）。"""
    docs = load_schema_docs(settings)
    head = docs.split("## 数据表介绍.docx")[0]
    physical = {p.stem for p in settings.parquet_dir.glob("*.parquet")}
    mapping: dict[str, str] = {}
    for m in re.finditer(r"([\u4e00-\u9fa5A-Za-z0-9（）()、]+?)\s*[（(]([A-Za-z_][A-Za-z0-9_]*)[)）]", head):
        cn, phys = m.group(1).strip(), m.group(2)
        if phys in physical:
            mapping[cn] = phys
    return mapping


def table_intros(settings: Settings) -> dict[str, str]:
    """介绍 docx 按 '## 中文名' 分节提取业务介绍，key 统一映射为物理表名。"""
    key = str(settings.table_data_dir)
    if key in _INTROS_CACHE:
        return _INTROS_CACHE[key]
    docs = load_schema_docs(settings)
    cn_map = _physical_name_map(settings)
    intros: dict[str, str] = {}
    current_name: str | None = None
    buf: list[str] = []

    def flush() -> None:
        nonlocal current_name, buf
        if current_name and buf:
            # 清单名与介绍标题存在 "X" / "X表" 两种形态，逐一尝试
            candidates = (current_name, current_name.rstrip("表"), current_name + "表")
            phys = next((cn_map[c] for c in candidates if c in cn_map), None)
            if phys:
                intros[phys] = "\n".join(buf).strip()[:400]
        buf = []

    for line in docs.split("\n"):
        if line.startswith("## ") and not line.startswith("## 数据"):
            flush()
            current_name = line[3:].strip()
        elif current_name is not None:
            buf.append(line)
    flush()
    logger.info("table_intros_built", extra={"tables": len(intros), "cn_map": len(cn_map)})
    _INTROS_CACHE[key] = intros
    return intros


def build_table_profiles(settings: Settings, tables: list[str] | None = None) -> str:
    """生成紧凑表画像：列名(中文描述/ dtype) + 1 行示例。tables=None 时全量。"""
    import pandas as pd

    colmap = column_dict(settings)
    paths = sorted(settings.parquet_dir.glob("*.parquet"))
    if tables is not None:
        wanted = set(tables)
        paths = [p for p in paths if p.stem in wanted]
    lines: list[str] = []
    for pq in paths:
        df = pd.read_parquet(pq)
        cols = ", ".join(
            f"{c}={colmap.get(c, '?')}({str(df[c].dtype)})" for c in df.columns
        )
        sample = df.head(1).to_csv(index=False)[:400]
        lines.append(f"### {pq.stem}（{len(df)} 行）\n列: {cols}\n示例: {sample}")
    return "\n\n".join(lines)
